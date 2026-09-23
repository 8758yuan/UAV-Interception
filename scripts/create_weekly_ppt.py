"""Generate a concise editable weekly report for the IBVS project."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / '汇报' / '无人机视觉拦截项目_本周进展汇报.pptx'
CHART = ROOT / '汇报' / 'latest_image_error.png'

W = Inches(13.333)
H = Inches(7.5)
FONT = 'Noto Sans CJK SC'

NAVY = RGBColor(15, 28, 48)
INK = RGBColor(29, 43, 60)
MUTED = RGBColor(100, 116, 139)
PAPER = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(18, 174, 184)
CYAN_DARK = RGBColor(10, 126, 137)
RED = RGBColor(224, 74, 74)
ORANGE = RGBColor(238, 148, 67)
GREEN = RGBColor(38, 166, 114)
LINE = RGBColor(218, 226, 235)
PALE_CYAN = RGBColor(229, 247, 248)
PALE_RED = RGBColor(253, 238, 238)
PALE_GREEN = RGBColor(232, 247, 239)


def add_shape(slide, kind, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind,
        x, y, w, h,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def set_run(run, size, color, bold=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(
    slide, text, x, y, w, h, size=20, color=INK, bold=False,
    align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    set_run(run, size, color, bold)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=18, gap=7):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)
    for index, (text, color, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        run = paragraph.add_run()
        run.text = text
        set_run(run, size, color, bold)
    return box


def add_header(slide, number, title, subtitle=None):
    add_text(slide, f'{number:02d}', Inches(0.55), Inches(0.35),
             Inches(0.55), Inches(0.4), 12, CYAN_DARK, True)
    add_text(slide, title, Inches(1.15), Inches(0.28), Inches(10.8),
             Inches(0.52), 27, NAVY, True)
    if subtitle:
        add_text(slide, subtitle, Inches(1.16), Inches(0.80), Inches(10.8),
                 Inches(0.35), 11.5, MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.17),
        Inches(12.2), Inches(0.025)
    )
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.color.rgb = LINE


def add_footer(slide, number):
    add_text(slide, 'IBVS 论文复现 · 周进展', Inches(0.58), Inches(7.12),
             Inches(3.3), Inches(0.22), 9, MUTED)
    add_text(slide, str(number), Inches(12.22), Inches(7.10),
             Inches(0.5), Inches(0.22), 9, MUTED, align=PP_ALIGN.RIGHT)


def add_arrow(slide, x1, y1, x2, y2, color=CYAN_DARK, width=2.2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_card(slide, x, y, w, h, tag, title, body, accent=CYAN):
    card = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, WHITE, LINE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.color.rgb = accent
    add_text(slide, tag, x + Inches(0.28), y + Inches(0.24),
             Inches(0.55), Inches(0.28), 11, accent, True)
    add_text(slide, title, x + Inches(0.28), y + Inches(0.58),
             w - Inches(0.52), Inches(0.42), 19, NAVY, True)
    add_text(slide, body, x + Inches(0.28), y + Inches(1.10),
             w - Inches(0.52), h - Inches(1.28), 13.5, MUTED)
    return card


def add_metric(slide, x, y, value, label, accent=CYAN):
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(2.18),
              Inches(1.12), WHITE, LINE)
    add_text(slide, value, x + Inches(0.16), y + Inches(0.16),
             Inches(1.86), Inches(0.48), 25, accent, True,
             align=PP_ALIGN.CENTER)
    add_text(slide, label, x + Inches(0.16), y + Inches(0.70),
             Inches(1.86), Inches(0.25), 10.5, MUTED,
             align=PP_ALIGN.CENTER)


def make_chart(path):
    times = [0.0, 0.503, 1.001, 1.5, 2.001, 2.5, 3.0, 3.5, 4.043]
    errors = [0.027, 0.135, 0.261, 0.362, 0.361, 0.288, 0.141, 0.098, 0.296]
    width, height = 1400, 620
    image = Image.new('RGB', (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 30)
        small = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 24)
        bold = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc', 31)
    except OSError:
        font = small = bold = ImageFont.load_default()
    left, top, right, bottom = 120, 75, 1330, 505
    for i in range(4):
        y = bottom - i * (bottom - top) / 3
        draw.line((left, y, right, y), fill=(224, 230, 237), width=2)
        draw.text((30, y - 16), f'{i * 0.2:.1f}', fill=(100, 116, 139), font=small)
    draw.line((left, top, left, bottom), fill=(55, 70, 88), width=3)
    draw.line((left, bottom, right, bottom), fill=(55, 70, 88), width=3)
    points = []
    for t, error in zip(times, errors):
        px = left + t / 4.1 * (right - left)
        py = bottom - error / 0.6 * (bottom - top)
        points.append((px, py))
    draw.line(points, fill=(18, 174, 184), width=7, joint='curve')
    for px, py in points:
        draw.ellipse((px - 8, py - 8, px + 8, py + 8), fill=(18, 174, 184))
    draw.text((left, 18), '主动拦截阶段图像误差', fill=(15, 28, 48), font=bold)
    draw.text((left, 545), '时间 / s', fill=(100, 116, 139), font=font)
    draw.text((right - 110, 545), '碰撞', fill=(224, 74, 74), font=font)
    image.save(path)


def build():
    make_chart(CHART)
    deck = Presentation()
    deck.slide_width = W
    deck.slide_height = H
    blank = deck.slide_layouts[6]

    # 1. Cover
    slide = deck.slides.add_slide(blank)
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    add_text(slide, '无人机视觉拦截', Inches(0.72), Inches(1.42),
             Inches(7.5), Inches(0.85), 34, WHITE, True)
    add_text(slide, '论文复现项目 · 本周进展汇报', Inches(0.76), Inches(2.28),
             Inches(7.5), Inches(0.55), 22, RGBColor(170, 231, 235), True)
    add_text(slide, 'ROS 2  ·  PX4  ·  Gazebo  ·  IBVS / DKF',
             Inches(0.77), Inches(3.18), Inches(6.8), Inches(0.42),
             14, RGBColor(185, 199, 216))
    add_text(slide, '汇报人  ________    ·    日期  2026.09.22',
             Inches(0.77), Inches(6.42), Inches(5.5), Inches(0.34),
             12, RGBColor(185, 199, 216))
    # Target and approach path motif.
    for diameter, color in ((2.3, RGBColor(54, 74, 98)),
                            (1.55, RGBColor(104, 55, 64)),
                            (0.82, RED)):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(9.75 + (2.3-diameter)/2),
            Inches(2.02 + (2.3-diameter)/2), Inches(diameter), Inches(diameter)
        )
        circle.fill.solid(); circle.fill.fore_color.rgb = color
        circle.line.color.rgb = color
    add_arrow(slide, Inches(7.75), Inches(4.65), Inches(10.25), Inches(3.47),
              RGBColor(104, 219, 225), 3.2)
    add_text(slide, '视觉闭环', Inches(7.55), Inches(4.78), Inches(2.2),
             Inches(0.36), 12, RGBColor(170, 231, 235), True)

    # 2. Reproduction flow
    slide = deck.slides.add_slide(blank)
    add_header(slide, 2, '论文复现总流程', '从图像观测到飞行器执行的完整闭环')
    labels = [
        ('01', '仿真场景', '无人机 / 气球'),
        ('02', '图像感知', '红色目标特征'),
        ('03', '状态观测', '延迟 DKF'),
        ('04', '论文控制律', 'IBVS + SO(3)'),
        ('05', 'PX4 执行', '角速度 + 推力'),
        ('06', '结果验证', '碰撞 / rosbag'),
    ]
    xs = [0.65, 2.75, 4.85, 6.95, 9.05, 11.15]
    for i, ((tag, title, body), x) in enumerate(zip(labels, xs)):
        fill = PALE_RED if i == 1 else (PALE_GREEN if i == 5 else WHITE)
        accent = RED if i == 1 else (GREEN if i == 5 else CYAN)
        add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.25),
                  Inches(1.55), Inches(2.05), fill, LINE)
        add_text(slide, tag, Inches(x + 0.15), Inches(2.43), Inches(0.42),
                 Inches(0.26), 10, accent, True)
        add_text(slide, title, Inches(x + 0.15), Inches(2.87), Inches(1.25),
                 Inches(0.62), 16, NAVY, True, align=PP_ALIGN.CENTER,
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, body, Inches(x + 0.15), Inches(3.55), Inches(1.25),
                 Inches(0.42), 11, MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            add_arrow(slide, Inches(x + 1.58), Inches(3.28),
                      Inches(x + 2.02), Inches(3.28), CYAN_DARK, 1.8)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.20), Inches(5.05),
              Inches(10.95), Inches(0.82), PALE_CYAN, PALE_CYAN)
    add_text(slide,
             '核心原则：控制器只使用机载 IMU、姿态与图像信息，不读取目标真实位置。',
             Inches(1.45), Inches(5.25), Inches(10.45), Inches(0.34),
             16, CYAN_DARK, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2)

    # 3. Weekly accomplishments
    slide = deck.slides.add_slide(blank)
    add_header(slide, 3, '本周完成内容', '完成静止目标闭环，并建立可重复测试入口')
    add_card(slide, Inches(0.68), Inches(1.55), Inches(5.85), Inches(2.12),
             'SIM', '仿真场景', '配置 PX4 + Gazebo；建立红色气球模型、接触检测与双视角界面。', RED)
    add_card(slide, Inches(6.80), Inches(1.55), Inches(5.85), Inches(2.12),
             'VISION', '视觉感知', '完成红色目标检测、归一化图像坐标与相机几何转换。', CYAN)
    add_card(slide, Inches(0.68), Inches(3.94), Inches(5.85), Inches(2.12),
             'CONTROL', '观测与控制', '实现延迟 DKF、论文 IBVS / SO(3) 控制律及 PX4 指令适配。', ORANGE)
    add_card(slide, Inches(6.80), Inches(3.94), Inches(5.85), Inches(2.12),
             'TEST', '试验与记录', '打通一键试验、rosbag 记录、安全门、碰撞确认和自动稳定流程。', GREEN)
    add_footer(slide, 3)

    # 4. Search and alignment
    slide = deck.slides.add_slide(blank)
    add_header(slide, 4, '关键改进：目标搜索与启动对准', '解决起飞后视野中无目标，以及启动偏移过大的问题')
    stages = [
        ('起飞悬停', '到达搜索高度'),
        ('无目标', '偏航 + 高度扫描'),
        ('发现目标', '立即停止扫描'),
        ('视觉对准', '调整偏航与高度'),
        ('稳定判定', '中心保持 0.8 s'),
        ('开始拦截', '切换角速度控制'),
    ]
    xs = [0.64, 2.74, 4.84, 6.94, 9.04, 11.14]
    for i, ((title, body), x) in enumerate(zip(stages, xs)):
        accent = RED if i == 2 else (GREEN if i == 5 else CYAN)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.48),
                                        Inches(2.02), Inches(0.56), Inches(0.56))
        circle.fill.solid(); circle.fill.fore_color.rgb = accent
        circle.line.color.rgb = accent
        add_text(slide, str(i + 1), Inches(x + 0.48), Inches(2.12),
                 Inches(0.56), Inches(0.25), 11, WHITE, True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, title, Inches(x), Inches(2.82), Inches(1.52),
                 Inches(0.40), 15, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x), Inches(3.30), Inches(1.52),
                 Inches(0.62), 10.5, MUTED, align=PP_ALIGN.CENTER)
        if i < 5:
            add_arrow(slide, Inches(x + 1.10), Inches(2.30),
                      Inches(x + 2.02), Inches(2.30), LINE, 2.0)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(1.03), Inches(4.73),
              Inches(11.28), Inches(1.10), PALE_CYAN, PALE_CYAN)
    add_text(slide, '当前关键参数', Inches(1.32), Inches(4.96),
             Inches(1.55), Inches(0.34), 14, CYAN_DARK, True)
    add_text(slide,
             '中心误差 ≤ 0.04   ·   释放阈值 0.08   ·   搜索偏航 0.20 rad/s   ·   垂直扫描 ±3 m',
             Inches(2.88), Inches(4.94), Inches(8.95), Inches(0.40),
             14, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4)

    # 5. Diagnosis and tuning
    slide = deck.slides.add_slide(blank)
    add_header(slide, 5, '测试问题定位与修正', '依据 rosbag 数据逐项定位，不直接盲目调参')
    headers = [('现象', 0.75, 2.45), ('定位结果', 3.35, 4.10), ('修正', 7.65, 4.92)]
    for title, x, w in headers:
        add_text(slide, title, Inches(x), Inches(1.48), Inches(w), Inches(0.38),
                 13, MUTED, True)
    rows = [
        ('远距离目标中途失败', '15 m 水平围栏提前触发降落', '安全范围调整为 45 m'),
        ('上升过程较慢', '搜索高度与目标高度差较大', '搜索高度提高至 10 m'),
        ('启动后目标大幅偏移', '角速度 0.6 rad/s，且允许 60° 视线偏差',
         '角速度 0.30 rad/s；视线边界 45°'),
    ]
    ys = [1.94, 3.25, 4.56]
    for index, ((symptom, cause, fix), y) in enumerate(zip(rows, ys)):
        fill = WHITE if index != 2 else PALE_RED
        add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(y),
                  Inches(12.0), Inches(1.02), fill, LINE)
        add_text(slide, symptom, Inches(0.91), Inches(y + 0.23),
                 Inches(2.15), Inches(0.50), 14, NAVY, True,
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, cause, Inches(3.38), Inches(y + 0.19),
                 Inches(3.96), Inches(0.58), 13, MUTED,
                 valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, fix, Inches(7.70), Inches(y + 0.19),
                 Inches(4.58), Inches(0.58), 13, GREEN, True,
                 valign=MSO_ANCHOR.MIDDLE)
    add_footer(slide, 5)

    # 6. Latest result
    slide = deck.slides.add_slide(blank)
    add_header(slide, 6, '最新静止目标测试结果',
               '数据来源：paper_static_20260922_210708')
    add_metric(slide, Inches(0.70), Inches(1.45), '0.027', '启动图像误差', CYAN_DARK)
    add_metric(slide, Inches(3.12), Inches(1.45), '4.05 s', '主动拦截时长', CYAN_DARK)
    add_metric(slide, Inches(5.54), Inches(1.45), '8.5 m/s', '最大速度', ORANGE)
    add_metric(slide, Inches(7.96), Inches(1.45), '27.4°', '最大倾角', ORANGE)
    add_metric(slide, Inches(10.38), Inches(1.45), '成功', '碰撞确认', GREEN)
    slide.shapes.add_picture(str(CHART), Inches(0.73), Inches(3.03),
                             width=Inches(7.50), height=Inches(3.30))
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(8.58), Inches(3.03),
              Inches(4.02), Inches(3.30), PALE_GREEN, PALE_GREEN)
    add_text(slide, '试验结论', Inches(8.93), Inches(3.38), Inches(3.32),
             Inches(0.42), 18, GREEN, True)
    add_rich_lines(slide, [
        ('✓  启动时目标接近图像中心', NAVY, True),
        ('✓  主动拦截约 4 秒完成', NAVY, True),
        ('✓  收到气球碰撞与变色确认', NAVY, True),
        ('当前状态：静止目标单次闭环验证完成', GREEN, True),
    ], Inches(8.92), Inches(4.00), Inches(3.34), Inches(1.85), 13.5, 10)
    add_footer(slide, 6)

    # 7. Next week
    slide = deck.slides.add_slide(blank)
    add_header(slide, 7, '当前结论与下一步', '从“单次可行”推进到“结果可重复、场景可扩展”')
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.60),
              Inches(4.00), Inches(4.65), PALE_GREEN, PALE_GREEN)
    add_text(slide, '本周结论', Inches(1.10), Inches(1.98), Inches(3.30),
             Inches(0.45), 21, GREEN, True)
    add_rich_lines(slide, [
        ('论文控制链路已完整运行', NAVY, True),
        ('视觉搜索与启动对准已加入', NAVY, True),
        ('静止气球实现成功撞击', NAVY, True),
        ('关键参数已集中管理', NAVY, True),
    ], Inches(1.10), Inches(2.76), Inches(3.20), Inches(2.60), 15.5, 13)
    add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(5.10), Inches(1.60),
              Inches(7.47), Inches(4.65), WHITE, LINE)
    add_text(slide, '下周计划', Inches(5.48), Inches(1.98), Inches(3.2),
             Inches(0.45), 21, NAVY, True)
    plans = [
        ('01', '静止目标重复试验', '统计成功率与误差峰值'),
        ('02', '参数敏感性分析', '比较深度、角速度与 Barrier 边界'),
        ('03', '移动目标验证', '接入论文目标运动模型'),
        ('04', '整理结果图表', '形成论文复现实验章节'),
    ]
    y = 2.68
    for tag, title, body in plans:
        add_text(slide, tag, Inches(5.50), Inches(y), Inches(0.48),
                 Inches(0.30), 11, CYAN_DARK, True)
        add_text(slide, title, Inches(6.05), Inches(y - 0.05), Inches(2.30),
                 Inches(0.38), 15, NAVY, True)
        add_text(slide, body, Inches(8.40), Inches(y - 0.03), Inches(3.65),
                 Inches(0.34), 12.5, MUTED)
        y += 0.78
    add_footer(slide, 7)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
